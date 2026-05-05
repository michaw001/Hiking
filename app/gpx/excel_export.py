from collections import defaultdict
from datetime import datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

MONTHS_DE = ["Januar", "Februar", "März", "April", "Mai", "Juni",
             "Juli", "August", "September", "Oktober", "November", "Dezember"]

COLOR_HEADER = "2E5FA3"
COLOR_TOTAL = "D9E4F5"
COLOR_ALT_ROW = "F2F6FC"


def style_header_row(ws, row, col_start, col_end, bg_color=COLOR_HEADER):
    fill = PatternFill("solid", fgColor=bg_color)
    font = Font(bold=True, color="FFFFFF")
    for col in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")


def style_total_row(ws, row, col_start, col_end):
    fill = PatternFill("solid", fgColor=COLOR_TOTAL)
    font = Font(bold=True)
    for col in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font


def set_border(ws, min_row, max_row, min_col, max_col):
    thin = Side(style="thin", color="AAAAAA")
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def write_all_tours_tab(wb, tours_dicts):
    ws = wb.create_sheet("Alle Touren")
    headers = ["Datum", "Dateiname", "Distanz (km)", "Dauer (min)", "Ø Speed (km/h)",
               "Aufstieg (m)", "Abstieg (m)", "Min Höhe (m)", "Max Höhe (m)", "Trackpoints"]
    ws.append(headers)
    style_header_row(ws, 1, 1, len(headers))

    for i, t in enumerate(sorted(tours_dicts, key=lambda x: x["date"] or datetime.min), start=2):
        row = [
            t["date"].strftime("%Y-%m-%d") if t["date"] else "",
            t["filename"],
            t["distance_km"],
            t["duration_min"],
            t["avg_speed_kmh"],
            t["ascent_m"],
            t["descent_m"],
            t["min_ele_m"],
            t["max_ele_m"],
            t["num_points"],
        ]
        ws.append(row)
        if i % 2 == 0:
            for col in range(1, len(headers) + 1):
                ws.cell(row=i, column=col).fill = PatternFill("solid", fgColor=COLOR_ALT_ROW)

    set_border(ws, 1, len(tours_dicts) + 1, 1, len(headers))
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 40
    for col in range(3, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16
    ws.freeze_panes = "A2"


def write_year_tab(wb, year, months_data):
    ws = wb.create_sheet(str(year))
    headers = ["Monat", "Touren", "Distanz (km)", "Dauer (h)", "Aufstieg (m)", "Ø km/Tour", "Ø Speed (km/h)"]
    ws.append(headers)
    style_header_row(ws, 1, 1, len(headers))

    data_rows = []
    for m_idx, month_name in enumerate(MONTHS_DE, start=1):
        d = months_data.get(m_idx, {})
        touren = d.get("count", 0)
        dist = round(d.get("distance_km", 0.0), 1)
        dur_h = round(d.get("duration_min", 0.0) / 60.0, 1) if d.get("duration_min") else 0.0
        aufstieg = round(d.get("ascent_m", 0.0), 0)
        avg_dist = round(dist / touren, 1) if touren > 0 else 0.0
        avg_speed_vals = d.get("speed_vals", [])
        avg_speed = round(sum(avg_speed_vals) / len(avg_speed_vals), 1) if avg_speed_vals else 0.0
        row = [month_name, touren, dist, dur_h, aufstieg, avg_dist, avg_speed]
        ws.append(row)
        data_rows.append(row)
        if m_idx % 2 == 0:
            for col in range(1, len(headers) + 1):
                ws.cell(row=m_idx + 1, column=col).fill = PatternFill("solid", fgColor=COLOR_ALT_ROW)

    total_row_idx = len(MONTHS_DE) + 2
    totals = ["Gesamt",
              sum(r[1] for r in data_rows),
              round(sum(r[2] for r in data_rows), 1),
              round(sum(r[3] for r in data_rows), 1),
              round(sum(r[4] for r in data_rows), 0),
              "", ""]
    ws.append(totals)
    style_total_row(ws, total_row_idx, 1, len(headers))
    set_border(ws, 1, total_row_idx, 1, len(headers))

    for col, width in zip(range(1, len(headers) + 1), [14, 8, 14, 10, 14, 12, 14]):
        ws.column_dimensions[get_column_letter(col)].width = width

    chart = BarChart()
    chart.type = "col"
    chart.title = f"Kilometer pro Monat {year}"
    chart.y_axis.title = "km"
    chart.x_axis.title = "Monat"
    chart.shape = 4
    chart.style = 10
    data_ref = Reference(ws, min_col=3, min_row=1, max_row=len(MONTHS_DE) + 1)
    cats = Reference(ws, min_col=1, min_row=2, max_row=len(MONTHS_DE) + 1)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats)
    chart.varyColors = True
    chart.width = 20
    chart.height = 12
    ws.add_chart(chart, f"A{total_row_idx + 2}")


def write_overview_tab(wb, yearly_data, years):
    ws = wb.create_sheet("Übersicht", 0)
    ws.sheet_view.showGridLines = False

    title_cell = ws["A1"]
    title_cell.value = "Wanderstatistik – Übersicht"
    title_cell.font = Font(bold=True, size=16, color=COLOR_HEADER)
    ws.merge_cells("A1:G1")
    ws.row_dimensions[1].height = 30

    ws.append([])
    headers = ["Jahr", "Touren", "Distanz (km)", "Dauer (h)", "Aufstieg (m)", "Ø km/Tour", "Ø Speed (km/h)"]
    ws.append(headers)
    style_header_row(ws, 3, 1, len(headers))

    for i, year in enumerate(sorted(years), start=4):
        d = yearly_data[year]
        touren = d["count"]
        dist = round(d["distance_km"], 1)
        dur_h = round(d["duration_min"] / 60.0, 1) if d["duration_min"] else 0.0
        aufstieg = round(d["ascent_m"], 0)
        avg_dist = round(dist / touren, 1) if touren > 0 else 0.0
        avg_speed_vals = d.get("speed_vals", [])
        avg_speed = round(sum(avg_speed_vals) / len(avg_speed_vals), 1) if avg_speed_vals else 0.0
        ws.append([year, touren, dist, dur_h, aufstieg, avg_dist, avg_speed])
        if i % 2 == 0:
            for col in range(1, len(headers) + 1):
                ws.cell(row=i, column=col).fill = PatternFill("solid", fgColor=COLOR_ALT_ROW)

    total_row = len(years) + 4
    all_d = yearly_data["_total"]
    tot_dist = round(all_d["distance_km"], 1)
    tot_dur_h = round(all_d["duration_min"] / 60.0, 1) if all_d["duration_min"] else 0.0
    tot_aufstieg = round(all_d["ascent_m"], 0)
    tot_avg = round(tot_dist / all_d["count"], 1) if all_d["count"] > 0 else 0.0
    tot_speed_vals = all_d.get("speed_vals", [])
    tot_avg_speed = round(sum(tot_speed_vals) / len(tot_speed_vals), 1) if tot_speed_vals else 0.0
    ws.append(["Gesamt", all_d["count"], tot_dist, tot_dur_h, tot_aufstieg, tot_avg, tot_avg_speed])
    style_total_row(ws, total_row, 1, len(headers))
    set_border(ws, 3, total_row, 1, len(headers))

    for col, width in zip(range(1, len(headers) + 1), [8, 8, 14, 10, 14, 12, 14]):
        ws.column_dimensions[get_column_letter(col)].width = width

    chart_start_row = total_row + 2
    n_years = len(years)

    chart_km = BarChart()
    chart_km.type = "col"
    chart_km.title = "Kilometer pro Jahr"
    chart_km.y_axis.title = "km"
    chart_km.style = 10
    data_km = Reference(ws, min_col=3, min_row=3, max_row=3 + n_years)
    cats = Reference(ws, min_col=1, min_row=4, max_row=3 + n_years)
    chart_km.add_data(data_km, titles_from_data=True)
    chart_km.set_categories(cats)
    chart_km.varyColors = True
    chart_km.width = 20
    chart_km.height = 14
    ws.add_chart(chart_km, f"A{chart_start_row}")

    chart_tours = BarChart()
    chart_tours.type = "col"
    chart_tours.title = "Touren pro Jahr"
    chart_tours.y_axis.title = "Anzahl"
    chart_tours.style = 10
    data_t = Reference(ws, min_col=2, min_row=3, max_row=3 + n_years)
    chart_tours.add_data(data_t, titles_from_data=True)
    chart_tours.set_categories(cats)
    chart_tours.varyColors = True
    chart_tours.width = 20
    chart_tours.height = 14
    ws.add_chart(chart_tours, f"I{chart_start_row}")

    chart_ascent = BarChart()
    chart_ascent.type = "col"
    chart_ascent.title = "Höhenmeter (Aufstieg) pro Jahr"
    chart_ascent.y_axis.title = "m"
    chart_ascent.style = 10
    data_a = Reference(ws, min_col=5, min_row=3, max_row=3 + n_years)
    chart_ascent.add_data(data_a, titles_from_data=True)
    chart_ascent.set_categories(cats)
    chart_ascent.varyColors = True
    chart_ascent.width = 20
    chart_ascent.height = 14
    ws.add_chart(chart_ascent, f"A{chart_start_row + 22}")


def generate_xlsx_for_user(tours):
    yearly_data = defaultdict(lambda: {"count": 0, "distance_km": 0.0, "duration_min": 0.0,
                                        "ascent_m": 0.0, "speed_vals": []})
    yearly_monthly = defaultdict(lambda: defaultdict(lambda: {"count": 0, "distance_km": 0.0,
                                                               "duration_min": 0.0, "ascent_m": 0.0,
                                                               "speed_vals": []}))
    tours_dicts = []
    for t in tours:
        d = t.to_dict()
        tours_dicts.append(d)
        if d["date"] is None:
            continue
        year = d["date"].year
        month = d["date"].month
        for key in ["_total", year]:
            yearly_data[key]["count"] += 1
            yearly_data[key]["distance_km"] += d["distance_km"]
            if d["duration_min"]:
                yearly_data[key]["duration_min"] += d["duration_min"]
            yearly_data[key]["ascent_m"] += d["ascent_m"]
            if d["avg_speed_kmh"]:
                yearly_data[key]["speed_vals"].append(d["avg_speed_kmh"])
        yearly_monthly[year][month]["count"] += 1
        yearly_monthly[year][month]["distance_km"] += d["distance_km"]
        if d["duration_min"]:
            yearly_monthly[year][month]["duration_min"] += d["duration_min"]
        yearly_monthly[year][month]["ascent_m"] += d["ascent_m"]
        if d["avg_speed_kmh"]:
            yearly_monthly[year][month]["speed_vals"].append(d["avg_speed_kmh"])

    years = sorted(k for k in yearly_data.keys() if k != "_total")

    wb = Workbook()
    wb.remove(wb.active)

    if years:
        write_overview_tab(wb, yearly_data, years)
        for year in years:
            write_year_tab(wb, year, yearly_monthly[year])

    write_all_tours_tab(wb, tours_dicts)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
